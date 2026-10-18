"""Import d'un classeur Excel de mots de passe (structure de mdps.xlsx).

Ne modifie jamais un compte existant : seules les lignes absentes sont ajoutées.
Ne journalise jamais de mot de passe.
"""
from __future__ import annotations

import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from app.images import normalize_image
from app.model import Account, GiftCard, Vault, new_id
from app.textutil import normalize

ACCOUNT_HEADERS = {"type": "category", "support": "name", "identifiant": "login", "mot de passe": "password"}
CARD_HEADERS = {"carte cadeaux": "name", "carte cadeau": "name", "code": "code", "pin": "pin",
                "montant restant": "remaining", "date fin": "end_date"}


def _general(v: float) -> str:
    """Affichage « Standard » d'Excel : 15 chiffres significatifs, virgule décimale."""
    if float(v).is_integer():
        return str(int(v))
    return f"{v:.15g}".replace(".", ",")


def _tokens(section: str) -> list[tuple[str, str]] | None:
    """Découpe une section de format en jetons ('ph', '0'|'#'|'?'), ('dot', '.'), ('lit', texte).

    Renvoie None si le format contient des éléments non gérés (dates, exposants…).
    """
    out, i = [], 0
    while i < len(section):
        ch = section[i]
        if ch == '"':
            j = section.find('"', i + 1)
            j = len(section) if j < 0 else j
            out.append(("lit", section[i + 1:j]))
            i = j + 1
            continue
        if ch == "\\" and i + 1 < len(section):
            out.append(("lit", section[i + 1]))
            i += 2
            continue
        if ch in "_*":  # espacement / remplissage : sans effet sur le texte
            i += 2
            continue
        if ch in "0#?":
            out.append(("ph", ch))
        elif ch == ".":
            out.append(("dot", "."))
        elif ch == ",":
            pass  # séparateur de milliers : ignoré (on garde les chiffres bruts)
        elif ch == "[":
            j = section.find("]", i)
            i = len(section) if j < 0 else j + 1
            continue
        elif ch.lower() in "dmyhse@":
            return None
        else:
            out.append(("lit", ch))
        i += 1
    return out


def _format_number(v: float, fmt: str) -> str:
    section = (fmt or "").split(";")[0].strip()
    if section in ("", "General", "Standard", "@"):
        return _general(v)
    toks = _tokens(section)
    if toks is None or not any(t == "ph" for t, _ in toks):
        return _general(v)
    if any(t == "lit" and "%" in x for t, x in toks):
        v = v * 100
    dot = next((i for i, (t, _) in enumerate(toks) if t == "dot"), len(toks))
    int_toks, dec_toks = toks[:dot], toks[dot + 1:]
    decimals = sum(1 for t, _ in dec_toks if t == "ph")
    neg = v < 0
    text = f"{abs(v):.{decimals}f}"
    int_digits, _, dec_digits = text.partition(".")
    if int_digits == "0" and any(c == "#" for t, c in int_toks if t == "ph") and not any(c == "0" for t, c in int_toks if t == "ph"):
        int_digits = ""
    # partie entière : placeholders remplis de droite à gauche, chiffres en trop devant le premier
    phs = [i for i, (t, _) in enumerate(int_toks) if t == "ph"]
    parts = [x for _, x in int_toks]
    digits = list(int_digits)
    for idx in reversed(phs):
        if digits:
            parts[idx] = digits.pop()
        else:
            parts[idx] = "0" if int_toks[idx][1] == "0" else (" " if int_toks[idx][1] == "?" else "")
    if digits and phs:
        parts[phs[0]] = "".join(digits) + parts[phs[0]]
    out = "".join(parts)
    if dec_toks:
        dparts, k = [], 0
        for t, x in dec_toks:
            if t == "ph":
                dparts.append(dec_digits[k] if k < len(dec_digits) else "0")
                k += 1
            else:
                dparts.append(x)
        out += "," + "".join(dparts)
    return ("-" if neg else "") + out


def cell_text(cell) -> str:
    """Texte tel qu'affiché dans Excel (formats numériques, zéros de tête, dates JJ/MM/AAAA).

    Les séparateurs de milliers sont ignorés : un code doit rester une suite de chiffres.
    """
    v = cell.value
    if v is None:
        return ""
    if isinstance(v, bool):
        return "VRAI" if v else "FAUX"
    if isinstance(v, (datetime, date)):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, (int, float)):
        return _format_number(v, cell.number_format)
    return str(v).strip()


def _find_header(ws, wanted: dict) -> tuple[int, dict[str, int]] | None:
    """Cherche dans les 10 premières lignes la ligne d'en-têtes ; renvoie (ligne, {champ: colonne})."""
    for row in ws.iter_rows(min_row=1, max_row=10):
        cols = {}
        for c in row:
            key = normalize(str(c.value)) if c.value is not None else ""
            if key in wanted and wanted[key] not in cols:
                cols[wanted[key]] = c.column
        if "name" in cols:
            return row[0].row, cols
    return None


def _clean_label(s: str) -> str:
    return re.sub(r"\s+", " ", s).strip()


def _merge_categories(raw: list[str]) -> dict[str, str]:
    """normalize(catégorie) → libellé retenu : le plus fréquent, à égalité celui avec seule l'initiale en majuscule."""
    groups: dict[str, Counter] = {}
    for c in raw:
        label = _clean_label(c)
        groups.setdefault(normalize(label), Counter())[label] += 1

    def rank(item):
        label, count = item
        sentence_case = label[:1].isupper() and label[1:] == label[1:].lower()
        return (count, sentence_case)

    return {k: max(cnt.items(), key=rank)[0] for k, cnt in groups.items()}


def _images_by_row(ws) -> dict[int, bytes]:
    """Images ancrées sur une ligne (1-based), seulement quand une seule image vise cette ligne."""
    rows: dict[int, list] = {}
    for img in getattr(ws, "_images", []):
        try:
            r = img.anchor._from.row + 1
        except AttributeError:
            continue
        rows.setdefault(r, []).append(img)
    out = {}
    for r, imgs in rows.items():
        if len(imgs) == 1:
            try:
                out[r] = imgs[0]._data()
            except Exception:  # image illisible : ignorée
                pass
    return out


def _account_key(category: str, name: str, login: str, password: str) -> tuple:
    return (normalize(category), normalize(name), normalize(login), password)


def import_xlsx(path: Path, vault: Vault, profile_id: str) -> tuple[int, int]:
    wb = load_workbook(path, data_only=True)
    added = skipped = 0

    ws = wb["Mots de passes"] if "Mots de passes" in wb.sheetnames else wb.worksheets[0]
    found = _find_header(ws, ACCOUNT_HEADERS)
    if found:
        hrow, cols = found
        rows = []
        for r in range(hrow + 1, ws.max_row + 1):
            rec = {f: cell_text(ws.cell(row=r, column=c)) for f, c in cols.items()}
            if rec.get("name"):
                rows.append((r, rec))
        labels = _merge_categories([rec.get("category", "") for _, rec in rows])
        images = _images_by_row(ws)
        existing = {_account_key(a.category, a.name, a.login, a.password) for a in vault.accounts}
        for r, rec in rows:
            category = labels.get(normalize(rec.get("category", "")), "")
            key = _account_key(category, rec["name"], rec.get("login", ""), rec.get("password", ""))
            if key in existing:
                skipped += 1
                continue
            image = None
            if r in images:
                try:
                    image = normalize_image(images[r])
                except ValueError:
                    image = None
            vault.accounts.append(Account(
                id=new_id(), name=rec["name"], category=category, login=rec.get("login", ""),
                password=rec.get("password", ""), image=image, profiles=[profile_id],
            ))
            existing.add(key)
            added += 1

    if "Cartes cadeaux" in wb.sheetnames:
        gs = wb["Cartes cadeaux"]
        found = _find_header(gs, CARD_HEADERS)
        if found:
            hrow, cols = found
            existing = {(normalize(g.name), g.code) for g in vault.giftcards}
            for r in range(hrow + 1, gs.max_row + 1):
                get = lambda f: gs.cell(row=r, column=cols[f]) if f in cols else None  # noqa: E731
                name = cell_text(get("name")) if get("name") else ""
                if not name:
                    continue
                code = cell_text(get("code")) if get("code") else ""
                if (normalize(name), code) in existing:
                    skipped += 1
                    continue
                remaining = 0.0
                if get("remaining") is not None and isinstance(get("remaining").value, (int, float)):
                    remaining = round(float(get("remaining").value), 2)
                end = get("end_date").value if get("end_date") is not None else None
                end_date = end.date().isoformat() if isinstance(end, datetime) else end.isoformat() if isinstance(end, date) else None
                vault.giftcards.append(GiftCard(
                    id=new_id(), name=name, code=code, pin=cell_text(get("pin")) if get("pin") else "",
                    initial=remaining or None, remaining=remaining, end_date=end_date, profiles=[profile_id],
                ))
                existing.add((normalize(name), code))
                added += 1
    return added, skipped
