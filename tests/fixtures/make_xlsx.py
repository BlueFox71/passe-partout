"""Génère un classeur factice reproduisant la structure de mdps.xlsx (aucune vraie donnée)."""
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from PIL import Image


def _png(color):
    buf = io.BytesIO()
    Image.new("RGB", (40, 40), color).save(buf, "PNG")
    buf.seek(0)
    return XlImage(buf)


ROWS = [  # Type, Support, Identifiant, Mot de passe
    ("Banque", "Boursorama", "jules.x", "Tr0mpette!Bleue"),
    ("Banque", "Code CB", "carte", 29017110),
    ("Divers", "Ligne5", "x", "y"),
    ("Etat", "Zero", "x", "06092597"),
    ("Jeu", "Flottant", "x", 133511.0),
    ("Réseau Social", "Facebook", "alexis@exemple.fr", "fb"),
    ("Réseau social", "Instagram", "alexis.photo", "insta"),
    ("Réseau social", "Discord", "jules_rp", "d20"),
    ("Plateforme  de jeu", "Steam", "julesgamer", "gaben"),
    ("Plateforme de jeu", "Epic", "julesgamer", "epic"),
    ("Etat", "Format", "x", 6092597),
    ("Santé", "Date", "x", datetime(2021, 3, 4)),
]


def make_xlsx(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Mots de passes"
    for col, h in zip("BCDE", ["Type", "Support", "Identifiant", "Mot de passe"]):
        ws[f"{col}2"] = h
    for i, row in enumerate(ROWS, start=3):
        for col, v in zip("BCDE", row):
            ws[f"{col}{i}"] = v
    ws["E13"].number_format = "00000000"
    ws["D8"].hyperlink = "mailto:alexis@exemple.fr"
    ws.cell(row=len(ROWS) + 5, column=3, value=None)  # lignes vides en fin de tableau
    ws.add_image(_png((255, 0, 0)), "F3")   # Boursorama : sans ambiguïté
    ws.add_image(_png((0, 255, 0)), "F5")   # Ligne5 : deux images → ambigu
    ws.add_image(_png((0, 0, 255)), "G5")
    gc = wb.create_sheet("Cartes cadeaux")
    for col, h in zip("BCDEF", ["Carte cadeaux", "Code", "Pin", "Montant restant", "Date fin"]):
        gc[f"{col}2"] = h
    for col, v in zip("BCDEF", ["Fnac", "ABC123", 4321, 25.5, datetime(2027, 1, 31)]):
        gc[f"{col}3"] = v
    wb.save(path)
    return path
